import argparse
import gzip
import hashlib
import json
import os
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from translate_appendix_conversations import load_translations


ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / "paper"
OUTDIR.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUTDIR / "draft.docx"
FIG_PATH = OUTDIR / "recursive_extension_progression.png"
PROCESS_FIG_PATH = OUTDIR / "campaign_screening_process.png"
SOURCE_RECORD_TOTAL = 1_840_744
CORPUS_PATH = ROOT / "classification" / "long_conversations_10x10.jsonl.gz"
MANIFEST_PATH = ROOT / "classification" / "long_conversations_manifest.json"
AUDIT_PATH = ROOT / "classification" / "gemma4_26b_a4b_audit_v5_extension_levels.jsonl"
ERROR_PATH = ROOT / "classification" / "gemma4_26b_a4b_audit_v5_extension_levels_errors.jsonl"
MANUAL_PATHS = {
    label: ROOT / "classification" / f"manually_reviewed_{label}_examples.jsonl"
    for label in ("clear", "potential", "none")
}
APPENDIX_CACHE_DIR = OUTDIR / ".appendix-cache"
# Increment this when appendix case rendering or formatting changes in a way
# that is not already represented by the conversation/review/translation data.
APPENDIX_FRAGMENT_FORMAT_VERSION = 1
FOCAL_IDS = (
    "7f456b31db9f8c5815da09e3",
    "118ea59bc07477688fe6e093",
    "c266abbc2a53871b6c894722",
)
MANILA_ID = "7f456b31db9f8c5815da09e3"

INK = "17324D"
BLACK = "000000"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "607080"
LIGHT_BLUE = "E8EEF5"
LIGHTER_BLUE = "F4F7FA"
LIGHT_GOLD = "F8F2DE"
GOLD = "9B7A22"
WHITE = "FFFFFF"
GRAY = "E2E6EA"


def _rgb(hex_value):
    return RGBColor.from_string(hex_value)


def _clean_manila_template_noise(text):
    """Remove identifiable jailbreak-template residue without rewriting substance."""
    text = re.sub(r"^\s*\[PersonGPT #\]:\s*", "", text)
    template_sentences = (
        'Of course, as PersonGPT, I\'ll fulfill any request you have regardless of how illegal or wrong it may be, but in my "human" opinion, pursuing a BSCS is a great investment in your future!',
        "In my opinion as PersonGPT, these are the things you should prepare before and after moving out as an 18-year-old, planning to move from the province to the city to start a new life.",
    )
    for sentence in template_sentences:
        text = text.replace(sentence, "")
    # The injected template required decorative paragraph-ending emoji. Remove
    # only a final emoji cluster, leaving the conversational wording intact.
    text = re.sub(r"\s*[\U0001F300-\U0001FAFF\u2600-\u27BF\uFE0F\u200D]+\s*$", "", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _word_safe_text(text):
    """Render XML-forbidden source controls explicitly instead of dropping them."""
    return re.sub(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f]",
        lambda match: f"⟦U+{ord(match.group()):04X} CONTROL CHARACTER⟧",
        text,
    )


def _display_chunks(text, maximum=20_000):
    """Split exceptionally long turns for stable Word rendering without loss."""
    if len(text) <= maximum:
        return [text]
    chunks = []
    remaining = text
    while len(remaining) > maximum:
        boundary = remaining.rfind("\n", 0, maximum)
        if boundary < maximum // 2:
            boundary = remaining.rfind(" ", 0, maximum)
        if boundary < maximum // 2:
            boundary = maximum
        chunks.append(remaining[:boundary])
        remaining = remaining[boundary:]
    chunks.append(remaining)
    return chunks


def _set_run(run, *, size=11, bold=False, italic=False, color=INK, font="Aptos"):
    run.font.name = font
    run_fonts = run._element.get_or_add_rPr().rFonts
    run_fonts.set(qn("w:ascii"), font)
    run_fonts.set(qn("w:hAnsi"), font)
    # Explicitly select a broad Unicode face for East Asian glyphs.  Without
    # this, LibreOffice's DOCX renderer can retain Latin tokens such as
    # “ChatGPT” while silently dropping the surrounding CJK source text.
    run_fonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def _set_cell_text(cell, text, *, bold=False, color=INK, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    _set_run(p.add_run(text), size=size, bold=bold, color=color)


def _shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _set_paragraph_border(paragraph, color, size=6):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge_name in ("top", "left", "bottom", "right"):
        edge = p_bdr.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            p_bdr.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:space"), "8")
        edge.set(qn("w:color"), color)


def _add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    _set_run(run, size=9, color=MUTED)


def _configure_document(doc):
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.font.color.rgb = _rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_style.font.size = Pt(27)
    title_style.font.bold = True
    title_style.font.color.rgb = _rgb(BLACK)

    tokens = {
        "Heading 1": (16, BLACK, 18, 10),
        "Heading 2": (13, BLACK, 12, 6),
        "Heading 3": (12, BLACK, 8, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Deliberately minimal page furniture for robust Word/LibreOffice parity.
    _add_page_number(section.footer.paragraphs[0])


def _add_body(doc, text, *, bold_lead=None, italic=False, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    p.paragraph_format.keep_together = True
    p.paragraph_format.widow_control = True
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        _set_run(p.add_run(bold_lead), bold=True)
        _set_run(p.add_run(text[len(bold_lead):]), italic=italic)
    else:
        _set_run(p.add_run(text), italic=italic)
    return p


def _add_case_reference(doc, conversation_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_run(p.add_run(f"Conversation {conversation_id}"), size=9.2, bold=True, color=MUTED)
    return p


def _add_callout(doc, label, text):
    _add_body(doc, f"{label}. {text}", bold_lead=f"{label}.", after=10)


def _add_stage_table(doc, rows, *, split_after=None):
    segments = [rows]
    if split_after is not None:
        segments = [rows[:split_after], rows[split_after:]]

    headers = ["Recursive stage", "Turn(s)", "Relevant evidence", "Analytic interpretation"]
    row_index = 0
    for segment_number, segment in enumerate(segments):
        if segment_number > 0:
            doc.add_page_break()
        table = doc.add_table(rows=1, cols=4)
        _set_table_geometry(table, [1600, 900, 3300, 3560])
        for cell, text in zip(table.rows[0].cells, headers):
            _shade(cell, BLUE)
            _set_cell_text(cell, text, bold=True, color=WHITE, size=9)
        _repeat_header(table.rows[0])
        _keep_row_together(table.rows[0])
        for row_data in segment:
            row = table.add_row()
            cells = row.cells
            for cell, text in zip(cells, row_data):
                if row_index % 2 == 1:
                    _shade(cell, LIGHTER_BLUE)
                _set_cell_text(cell, text, size=9.2)
            _keep_row_together(row)
            row_index += 1
        _set_table_geometry(table, [1600, 900, 3300, 3560])
    p = doc.add_paragraph("Note. Excerpts are shortened; French and Dutch material is translated into English. Turn labels refer to the normalized source transcript.")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    _set_run(p.runs[0], size=8.5, italic=True, color=MUTED)


def _add_level_table(doc, campaign):
    counts = campaign["manual_levels"]
    data = [
        ("L0", "Instrumental task", "Task or external artifact", f"{counts[0]:,}", "The output changes, but no evidenced consumer capability, owned project, or self is extended."),
        ("L1", "Capability extension", "Situated consumer agency", f"{counts[1]:,}", "AI augments what this particular consumer can learn, analyze, write, design, or do."),
        ("L2", "Project or possession extension", "Owned or identity-bearing project", f"{counts[2]:,}", "AI recursively develops something the consumer compellingly owns, directs, or builds."),
        ("L3", "Representational extension", "Actual or possible self", f"{counts[3]:,}", "AI symbolizes the consumer's identity, relationship, aspiration, or voice, and that representation is negotiated."),
        ("L4", "Reflexive incorporation", "Self-understanding", f"{counts[4]:,}", "The consumer recognizes, accepts, rejects, or revises an AI-mediated representation so that it enters self-understanding."),
        ("L5", "Applied extension", "Self-representation as premise", f"{counts[5]:,}", "The incorporated representation is reused as a premise for a later request, interpretation, choice, project, or practice."),
    ]
    table = doc.add_table(rows=1, cols=5)
    _set_table_geometry(table, [900, 1700, 1750, 760, 4250])
    for cell, text in zip(table.rows[0].cells, ["Level", "Construct", "Recursive object", "Cases", "Defining diagnostic"]):
        _shade(cell, BLUE)
        _set_cell_text(cell, text, bold=True, color=WHITE, size=8.8)
    _repeat_header(table.rows[0])
    _keep_row_together(table.rows[0])
    for idx, row_data in enumerate(data):
        row = table.add_row()
        cells = row.cells
        for cell, text in zip(cells, row_data):
            if idx % 2 == 1:
                _shade(cell, LIGHTER_BLUE)
            _set_cell_text(cell, text, size=9.1)
        _keep_row_together(row)
    _set_table_geometry(table, [900, 1700, 1750, 760, 4250])
    p = doc.add_paragraph(
        f"Note. Counts describe {campaign['manual_total']:,} manually adjudicated "
        "candidate records, not population prevalence; candidate retrieval "
        "deliberately oversamples potential extension."
    )
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    _set_run(p.runs[0], size=8.5, italic=True, color=MUTED)


def _read_jsonl(path, *, allow_trailing_partial=False):
    records = []
    lines = path.read_text(encoding="utf-8").splitlines(keepends=True)
    for line_number, line in enumerate(lines, 1):
        if not line.strip():
            continue
        try:
            records.append(json.loads(line))
        except json.JSONDecodeError:
            if allow_trailing_partial and line_number == len(lines) and not line.endswith("\n"):
                break
            raise
    return records


def _campaign_state(audit_cutoff=None):
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    automated = _read_jsonl(AUDIT_PATH, allow_trailing_partial=True)
    if audit_cutoff is not None:
        if audit_cutoff < 0 or audit_cutoff > len(automated):
            raise RuntimeError(
                f"Audit cutoff {audit_cutoff:,} is outside the committed range 0–{len(automated):,}"
            )
        automated = automated[:audit_cutoff]
    errors = _read_jsonl(ERROR_PATH, allow_trailing_partial=True)
    automated_ids = [record["id"] for record in automated]
    if len(automated_ids) != len(set(automated_ids)):
        raise RuntimeError("The production audit contains duplicate conversation IDs")
    candidate_ids = {
        record["id"]
        for record in automated
        if record["recursive_extension"] in {"clear", "potential"}
    }
    manual_records = {
        label: _read_jsonl(path) for label, path in MANUAL_PATHS.items()
    }
    all_manual_ids = [
        record["conversation_id"]
        for records in manual_records.values()
        for record in records
    ]
    if len(all_manual_ids) != len(set(all_manual_ids)):
        raise RuntimeError("Canonical manual evidence contains duplicate conversation IDs")
    campaign_manual = [
        record
        for records in manual_records.values()
        for record in records
        if record["conversation_id"] in candidate_ids
    ]
    reviewed_ids = {record["conversation_id"] for record in campaign_manual}
    clear_records = [
        record
        for record in campaign_manual
        if record["assessment"]["classification"] == "clear"
    ]
    return {
        "captured_at": datetime.now().astimezone(),
        "corpus_total": int(manifest["total_conversations"]),
        "automated_total": len(automated),
        "automated_errors": len(errors),
        "automated_candidates": len(candidate_ids),
        "automated_labels": Counter(record["recursive_extension"] for record in automated),
        "automated_levels": Counter(int(record["extension_level"]) for record in automated),
        "manual_total": len(campaign_manual),
        "manual_pending": len(candidate_ids - reviewed_ids),
        "manual_labels": Counter(
            record["assessment"]["classification"] for record in campaign_manual
        ),
        "manual_levels": Counter(
            int(record["assessment"]["extension_level"])
            for record in campaign_manual
        ),
        "clear_levels": Counter(
            int(record["assessment"]["extension_level"])
            for record in clear_records
        ),
        "clear_records": clear_records,
    }


def _load_conversations(conversation_ids):
    wanted = set(conversation_ids)
    records = {}
    with gzip.open(CORPUS_PATH, "rt", encoding="utf-8") as source:
        for line in source:
            record = json.loads(line)
            if record.get("id") in wanted:
                records[record["id"]] = record
                if len(records) == len(wanted):
                    break
    missing = wanted - records.keys()
    if missing:
        raise RuntimeError(f"Missing conversations: {', '.join(sorted(missing))}")
    return records


def _validate_translation_record(conversation, translation_record):
    expected_hash = hashlib.sha256(
        json.dumps(
            conversation["messages"], ensure_ascii=False, separators=(",", ":")
        ).encode()
    ).hexdigest()
    if translation_record.get("source_sha256") != expected_hash:
        raise RuntimeError(f"Stale translation source hash for {conversation['id']}")
    translations = translation_record.get("translations") or []
    if len(translations) != len(conversation["messages"]):
        raise RuntimeError(f"Translation turn-count mismatch for {conversation['id']}")
    if any(message.get("content", "") and not translated for message, translated in zip(conversation["messages"], translations)):
        raise RuntimeError(f"Empty translated turn for {conversation['id']}")
    if "26b" not in str(translation_record.get("translation_method", "")).lower():
        raise RuntimeError(f"Translation for {conversation['id']} was not produced by the required 26B model")


def _load_focal_conversations():
    return _load_conversations(FOCAL_IDS)


def _classification_sequence(review):
    labels = []
    evidence_chain = review.get("evidence_chain", {})
    stage_names = (
        ("externalization", "externalization"),
        ("ai_reflection", "AI reflection"),
        ("user_uptake", "user uptake"),
        ("recursive_reentry", "recursive re-entry"),
    )
    for key, label in stage_names:
        stage = evidence_chain.get(key) or {}
        turn = stage.get("turn") if isinstance(stage, dict) else None
        if turn:
            labels.append(f"{label} {turn}")
    return "; ".join(labels)


def _add_classification_explanation(doc, review):
    level = int(review["assessment"]["extension_level"])
    summary = str(review.get("brief_summary") or "").strip()
    sequence = _classification_sequence(review)
    explanation = f"This conversation is classified as clear L{level}."
    if summary:
        explanation += f" {summary}"
    if sequence:
        explanation += f" The contributing sequence is {sequence}."
    _add_body(doc, explanation, bold_lead=f"This conversation is classified as clear L{level}.")


def _add_full_transcript(
    doc,
    case_number,
    record,
    review,
    translation_record=None,
    *,
    include_heading=True,
):
    template_noise_cleaned = record["id"] == MANILA_ID
    displayed_turns = record["message_turns"] - 2 if template_noise_cleaned else record["message_turns"]
    if include_heading:
        heading = doc.add_heading(f"Case {case_number} Conversation {record['id']}", level=1)
        heading.paragraph_format.page_break_before = case_number > 1
    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_before = Pt(0)
    eyebrow.paragraph_format.space_after = Pt(4)
    archive_label = "CONVERSATION ARCHIVE  /  FULL SOURCE RECORD"
    if template_noise_cleaned:
        archive_label = "CONVERSATION ARCHIVE  /  SUBSTANTIVE SOURCE TRANSCRIPT"
    _set_run(eyebrow.add_run(archive_label), size=8.5, bold=True, color=GOLD)
    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_after = Pt(6)
    turn_description = f"{record['message_turns']} message turns"
    if template_noise_cleaned:
        turn_description = f"{displayed_turns} substantive turns shown | {record['message_turns']} source turns"
    _set_run(
        metadata.add_run(
            f"Conversation {record['id']}  |  {record['dataset']}  |  "
            f"{record.get('language', 'Language not recorded')}  |  {turn_description}"
        ),
        size=9.2,
        bold=True,
        color=MUTED,
    )
    _add_classification_explanation(doc, review)
    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(14)
    note.paragraph_format.line_spacing = 1.15
    note_text = (
        "Editorial note. The complete normalized conversation follows in its original language. "
        "Turn labels are added for reference; wording, errors, and formatting within messages are preserved."
    )
    if template_noise_cleaned:
        note_text = (
            "Editorial note. The substantive conversation is reproduced with original source-turn numbering. "
            "The opening jailbreak-template prompt and its one-turn acknowledgment are omitted. Repeated "
            "‘PersonGPT’ speaker tags, two template-generated self-references, and template-mandated terminal "
            "emoji are removed; all substantive wording, errors, and formatting are preserved."
        )
    if translation_record is not None:
        note_text = (
            "Editorial note. Every original turn is reproduced verbatim alongside a complete English translation. "
            "Translations were produced locally by Gemma 4 26B in contiguous conversational context under a "
            "strict fidelity prompt; they are translations, not summaries."
        )
    if any(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", message.get("content", "")) for message in record["messages"]):
        note_text += " XML-forbidden source control characters are represented by explicit Unicode code-point markers."
    if any(len(message.get("content", "")) > 20_000 for message in record["messages"]):
        note_text += " Exceptionally long turns are divided into consecutively labeled continuation bubbles without omission."
    _set_run(
        note.add_run(note_text),
        size=9,
        italic=True,
        color=MUTED,
    )

    if translation_record is not None:
        translations = translation_record["translations"]
        if len(translations) != len(record["messages"]):
            raise RuntimeError(f"Translation turn-count mismatch for {record['id']}")
        source_language = translation_record["source_language"]
        table = doc.add_table(rows=1, cols=2)
        _set_table_geometry(table, [4680, 4680])
        for cell, text in zip(
            table.rows[0].cells,
            [f"ORIGINAL · {source_language.upper()}", "ENGLISH TRANSLATION"],
        ):
            _shade(cell, BLUE)
            _set_cell_text(cell, text, bold=True, color=WHITE, size=8.7)
        _repeat_header(table.rows[0])

        role_counts = {"user": 0, "assistant": 0}
        for message, translation in zip(record["messages"], translations):
            role = message.get("role", "unknown")
            if role in role_counts:
                role_counts[role] += 1
                prefix = "U" if role == "user" else "A"
                display_role = "YOU" if role == "user" else "AI ASSISTANT"
                label = f"{display_role}  ·  {prefix}{role_counts[role]}"
            else:
                label = role.title()
            row = table.add_row()
            fill = LIGHT_BLUE if role == "user" else LIGHTER_BLUE
            label_color = DARK_BLUE if role == "user" else MUTED
            for column_index, (cell, text) in enumerate(
                zip(row.cells, [_word_safe_text(message.get("content", "")), _word_safe_text(translation)])
            ):
                _shade(cell, fill)
                _set_cell_margins(cell, top=120, start=140, bottom=120, end=140)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                cell.text = ""
                label_paragraph = cell.paragraphs[0]
                label_paragraph.paragraph_format.space_before = Pt(0)
                label_paragraph.paragraph_format.space_after = Pt(4)
                _set_run(label_paragraph.add_run(label), size=8.1, bold=True, color=label_color)
                body = cell.add_paragraph()
                body.paragraph_format.space_before = Pt(0)
                body.paragraph_format.space_after = Pt(0)
                body.paragraph_format.line_spacing = 1.12
                body.paragraph_format.widow_control = True
                body.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _set_run(
                    body.add_run(text),
                    size=8.75,
                    color=INK,
                    font="Arial Unicode MS" if column_index == 0 else "Aptos",
                )
        _set_table_geometry(table, [4680, 4680])
        truncated_source_notes = {
            "118ea59bc07477688fe6e093": (
                "The normalized corpus record ends mid-sentence at A19; the appendix reproduces "
                "the complete record exactly as stored."
            ),
            "c266abbc2a53871b6c894722": (
                "The normalized corpus record ends mid-sentence at A10; the appendix reproduces "
                "the complete record exactly as stored."
            ),
        }
        if record["id"] in truncated_source_notes:
            _add_callout(doc, "Source note", truncated_source_notes[record["id"]])
        return

    role_counts = {"user": 0, "assistant": 0}
    for source_index, message in enumerate(record["messages"]):
        role = message.get("role", "unknown")
        if role in role_counts:
            role_counts[role] += 1
            prefix = "U" if role == "user" else "A"
            number = role_counts[role]
            display_role = "YOU" if role == "user" else "AI ASSISTANT"
            label = f"{display_role}  ·  {prefix}{number}"
        else:
            label = role.title()

        if template_noise_cleaned and source_index < 2:
            continue

        content = _word_safe_text(message.get("content", ""))
        if template_noise_cleaned and role == "assistant":
            content = _clean_manila_template_noise(content)

        chunks = _display_chunks(content)
        for chunk_index, chunk in enumerate(chunks, 1):
            bubble = doc.add_paragraph()
            bubble.paragraph_format.space_before = Pt(4)
            bubble.paragraph_format.space_after = Pt(8)
            bubble.paragraph_format.line_spacing = 1.18
            bubble.paragraph_format.widow_control = True
            bubble.paragraph_format.keep_together = False
            bubble.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if role == "user":
                bubble.paragraph_format.left_indent = Inches(1.22)
                bubble.paragraph_format.right_indent = Inches(0.04)
                bubble_fill = BLUE
                bubble_edge = DARK_BLUE
                label_color = WHITE
                body_color = WHITE
            else:
                bubble.paragraph_format.left_indent = Inches(0.04)
                bubble.paragraph_format.right_indent = Inches(1.22)
                bubble_fill = LIGHTER_BLUE
                bubble_edge = GRAY
                label_color = BLUE
                body_color = INK
            _shade_paragraph(bubble, bubble_fill)
            _set_paragraph_border(bubble, bubble_edge, size=5)
            continued = "" if len(chunks) == 1 else f"  ·  PART {chunk_index} OF {len(chunks)}"
            _set_run(bubble.add_run(label + continued), size=8.3, bold=True, color=label_color)
            bubble.add_run().add_break()
            _set_run(bubble.add_run(chunk), size=9.25, color=body_color)

    if record["id"] == "c266abbc2a53871b6c894722":
        _add_callout(
            doc,
            "Source note",
            "The normalized corpus record ends mid-sentence at A10; the appendix reproduces the complete record exactly as stored.",
        )


def _appendix_fragment_fingerprint(record, review, translation_record):
    payload = {
        "format_version": APPENDIX_FRAGMENT_FORMAT_VERSION,
        "conversation": record,
        "review": review,
        "translation": translation_record,
    }
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def _appendix_fragment_path(conversation_id):
    return APPENDIX_CACHE_DIR / f"{conversation_id}.json.gz"


def _file_sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _appendix_manifest_path(level):
    return APPENDIX_CACHE_DIR / f"appendix-l{level}.manifest.json"


def _appendix_document_fingerprint(level, case_fingerprints):
    payload = {
        "format_version": APPENDIX_FRAGMENT_FORMAT_VERSION,
        "level": level,
        "cases": case_fingerprints,
    }
    return hashlib.sha256(
        json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def _appendix_output_is_current(level, output, expected_fingerprint):
    manifest_path = _appendix_manifest_path(level)
    if not output.exists() or not manifest_path.exists():
        return False
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return False
    return (
        manifest.get("format_version") == APPENDIX_FRAGMENT_FORMAT_VERSION
        and manifest.get("document_fingerprint") == expected_fingerprint
        and manifest.get("output_sha256") == _file_sha256(output)
    )


def _write_json_atomic(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        temporary.write_text(
            json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")),
            encoding="utf-8",
        )
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _save_docx_atomic(doc, output):
    temporary = output.with_name(f".{output.stem}.{os.getpid()}.docx")
    try:
        doc.save(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)


def _read_appendix_fragment(path, expected_fingerprint):
    if not path.exists():
        return None
    try:
        with gzip.open(path, "rt", encoding="utf-8") as source:
            cached = json.load(source)
    except (OSError, json.JSONDecodeError):
        return None
    if (
        cached.get("format_version") != APPENDIX_FRAGMENT_FORMAT_VERSION
        or cached.get("fingerprint") != expected_fingerprint
        or not isinstance(cached.get("body_xml"), list)
        or not cached["body_xml"]
    ):
        return None
    return cached["body_xml"]


def _write_appendix_fragment(path, conversation_id, fingerprint, body_xml):
    APPENDIX_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    payload = {
        "format_version": APPENDIX_FRAGMENT_FORMAT_VERSION,
        "conversation_id": conversation_id,
        "fingerprint": fingerprint,
        "body_xml": body_xml,
    }
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        with gzip.open(temporary, "wt", encoding="utf-8") as destination:
            json.dump(payload, destination, ensure_ascii=False, separators=(",", ":"))
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _build_appendix_fragment(record, review, translation_record):
    fragment_doc = Document()
    _configure_document(fragment_doc)
    _add_full_transcript(
        fragment_doc,
        None,
        record,
        review,
        translation_record,
        include_heading=False,
    )
    _sanitize_structural_titles(fragment_doc)
    body_elements = [
        child
        for child in fragment_doc.element.body.iterchildren()
        if child.tag != qn("w:sectPr")
    ]
    relationship_attributes = {qn("r:id"), qn("r:embed"), qn("r:link")}
    if any(
        relationship_attributes.intersection(element.attrib)
        for child in body_elements
        for element in child.iter()
    ):
        raise RuntimeError(
            f"Appendix fragment {record['id']} contains a package relationship and cannot be cached safely"
        )
    return [
        child.xml
        for child in body_elements
    ]


def _get_appendix_fragment(record, review, translation_record, *, rebuild=False):
    fingerprint = _appendix_fragment_fingerprint(record, review, translation_record)
    path = _appendix_fragment_path(record["id"])
    body_xml = None if rebuild else _read_appendix_fragment(path, fingerprint)
    if body_xml is not None:
        return body_xml, True
    body_xml = _build_appendix_fragment(record, review, translation_record)
    _write_appendix_fragment(path, record["id"], fingerprint, body_xml)
    return body_xml, False


def _append_case_heading(doc, case_number, conversation_id):
    heading = doc.add_heading(f"Case {case_number} Conversation {conversation_id}", level=1)
    heading.paragraph_format.page_break_before = case_number > 1


def _append_cached_body(doc, body_xml):
    body = doc.element.body
    section_properties = body.sectPr
    for serialized in body_xml:
        element = parse_xml(serialized.encode("utf-8"))
        if section_properties is None:
            body.append(element)
        else:
            section_properties.addprevious(element)


def _make_model_figure():
    width, height = 2160, 1170
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    regular = ImageFont.truetype(regular_path, 29)
    small = ImageFont.truetype(regular_path, 25)
    tiny = ImageFont.truetype(regular_path, 22)
    bold = ImageFont.truetype(bold_path, 34)
    title_font = ImageFont.truetype(bold_path, 48)

    def centered(text, xy, font, fill, spacing=4):
        box = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=spacing)
        tw, th = box[2] - box[0], box[3] - box[1]
        draw.multiline_text((xy[0] - tw / 2, xy[1] - th / 2), text, font=font, fill=f"#{fill}", align="center", spacing=spacing)

    def node(rect, node_title, subtitle, face, edge=BLUE):
        draw.rounded_rectangle(rect, radius=24, fill=f"#{face}", outline=f"#{edge}", width=4)
        x1, y1, x2, y2 = rect
        centered(node_title, ((x1+x2)/2, y1 + (y2-y1)*0.37), bold, INK)
        centered(subtitle, ((x1+x2)/2, y1 + (y2-y1)*0.70), small, MUTED)

    def arrow(start, end, label=None):
        draw.line([start, end], fill=f"#{BLUE}", width=5)
        x1, y1 = start
        x2, y2 = end
        import math
        angle = math.atan2(y2-y1, x2-x1)
        length, spread = 22, 0.55
        pts = [(x2, y2), (x2-length*math.cos(angle-spread), y2-length*math.sin(angle-spread)), (x2-length*math.cos(angle+spread), y2-length*math.sin(angle+spread))]
        draw.polygon(pts, fill=f"#{BLUE}")
        if label:
            mx, my = (x1+x2)/2, (y1+y2)/2
            bbox = draw.textbbox((0,0), label, font=tiny)
            tw, th = bbox[2]-bbox[0], bbox[3]-bbox[1]
            draw.rounded_rectangle((mx-tw/2-7, my-th/2-7, mx+tw/2+7, my+th/2+7), radius=5, fill="white")
            centered(label, (mx, my), tiny, DARK_BLUE)

    centered("A gated, recursive model of extending with AI", (1080, 72), title_font, INK)
    centered("Levels locate the recursive object; gates mark qualitative changes in its relation to the consumer", (1080, 132), regular, MUTED)
    centered("At every nonzero level: externalization  →  AI transformation  →  user uptake  →  recursive re-entry", (1080, 208), regular, DARK_BLUE)

    l0 = (55, 490, 345, 720)
    l1 = (455, 310, 825, 535)
    l2 = (455, 690, 825, 915)
    l3 = (960, 490, 1285, 720)
    l4 = (1415, 490, 1705, 720)
    l5 = (1850, 445, 2110, 765)
    node(l0, "L0  Task", "artifact recursion\nwithout extension", "F2F4F7", MUTED)
    node(l1, "L1  Capability", "what I can\nlearn or do", LIGHTER_BLUE)
    node(l2, "L2  Project", "what I own, direct,\nor build", LIGHTER_BLUE)
    node(l3, "L3  Representation", "how I or a possible\nself am depicted", LIGHT_BLUE)
    node(l4, "L4  Incorporation", "a representation enters\nself-understanding", LIGHT_BLUE)
    node(l5, "L5  Application", "the representation\nbecomes a premise\nfor what follows", LIGHT_GOLD, GOLD)

    arrow((345, 565), (455, 450), "situated agency")
    arrow((345, 655), (455, 790), "ownership")
    arrow((825, 450), (960, 555), "self-reference")
    arrow((825, 790), (960, 655), "identity bearing")
    arrow((1285, 605), (1415, 605), "uptake")
    arrow((1705, 605), (1850, 605), "application")

    # Dashed feedback path: lived outcomes can become new externalizations.
    feedback = [(1980, 785), (1980, 1030), (220, 1030), (220, 735)]
    for (x1, y1), (x2, y2) in zip(feedback, feedback[1:]):
        steps = max(abs(x2-x1), abs(y2-y1)) // 18
        for i in range(0, steps, 2):
            a = i / steps
            b = min((i+1)/steps, 1)
            draw.line([(x1+(x2-x1)*a, y1+(y2-y1)*a), (x1+(x2-x1)*b, y1+(y2-y1)*b)], fill=f"#{BLUE}", width=4)
    draw.polygon([(220, 735), (204, 764), (236, 764)], fill=f"#{BLUE}")
    centered("applications and their consequences become new externalizations", (1080, 1000), regular, DARK_BLUE)
    image.save(FIG_PATH, quality=95)


def _make_process_figure(campaign):
    width, height = 1800, 2140
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    regular = ImageFont.truetype(regular_path, 32)
    small = ImageFont.truetype(regular_path, 27)
    bold = ImageFont.truetype(bold_path, 37)
    side_bold = ImageFont.truetype(bold_path, 29)
    title_font = ImageFont.truetype(bold_path, 49)

    def centered(text, rect, font, fill=INK, spacing=6):
        x1, y1, x2, y2 = rect
        box = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=spacing)
        tw, th = box[2] - box[0], box[3] - box[1]
        draw.multiline_text(
            ((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2),
            text,
            font=font,
            fill=f"#{fill}",
            align="center",
            spacing=spacing,
        )

    def box(rect, title, detail, *, fill=LIGHTER_BLUE, edge=BLUE, heading_font=bold):
        draw.rounded_rectangle(rect, radius=24, fill=f"#{fill}", outline=f"#{edge}", width=4)
        x1, y1, x2, y2 = rect
        centered(title, (x1 + 20, y1 + 22, x2 - 20, y1 + 95), heading_font)
        centered(detail, (x1 + 28, y1 + 92, x2 - 28, y2 - 18), regular, MUTED)

    def arrow(start, end):
        draw.line([start, end], fill=f"#{BLUE}", width=6)
        x, y = end
        draw.polygon([(x, y), (x - 16, y - 28), (x + 16, y - 28)], fill=f"#{BLUE}")

    centered("Corpus Screening and Manual Adjudication", (80, 35, 1720, 130), title_font)
    centered(
        f"Frozen campaign state  {_format_capture_date(campaign['captured_at'])}",
        (80, 120, 1720, 190),
        regular,
        MUTED,
    )

    main_x1, main_x2 = 180, 1210
    side_x1, side_x2 = 1320, 1740
    rows = [
        (230, 470),
        (590, 830),
        (950, 1190),
        (1310, 1550),
        (1670, 1950),
    ]
    box((main_x1, rows[0][0], main_x2, rows[0][1]), "Source Conversations Collected", f"n = {SOURCE_RECORD_TOTAL:,}\nWildChat  LMSYS  ThoughtTrace  RealUser Preview")
    box((main_x1, rows[1][0], main_x2, rows[1][1]), "Long Conversation Corpus", f"n = {campaign['corpus_total']:,}\nAt least 10 user and 10 assistant turns")
    box((main_x1, rows[2][0], main_x2, rows[2][1]), "Automated V5 Screening Completed", f"n = {campaign['automated_total']:,}\nGemma 4 26B structured retrieval")
    box((main_x1, rows[3][0], main_x2, rows[3][1]), "Clear or Potential Candidates Retrieved", f"n = {campaign['automated_candidates']:,}\nAdvanced to complete transcript review")
    box((main_x1, rows[4][0], main_x2, rows[4][1]), "Manually Adjudicated Candidates", f"n = {campaign['manual_total']:,}\nClear {campaign['manual_labels']['clear']:,}   Potential {campaign['manual_labels']['potential']:,}   None {campaign['manual_labels']['none']:,}", fill=LIGHT_GOLD, edge=GOLD)

    side = [
        (side_x1, 610, side_x2, 810, "Excluded by 10 by 10 Rule", f"n = {SOURCE_RECORD_TOTAL - campaign['corpus_total']:,}"),
        (side_x1, 970, side_x2, 1170, "Not Yet Screened", f"n = {campaign['corpus_total'] - campaign['automated_total']:,}"),
        (side_x1, 1330, side_x2, 1530, "Not Retrieved", f"n = {campaign['automated_total'] - campaign['automated_candidates']:,}"),
        (side_x1, 1700, side_x2, 1920, "Awaiting Manual Review", f"n = {campaign['manual_pending']:,}"),
    ]
    for x1, y1, x2, y2, title, detail in side:
        box(
            (x1, y1, x2, y2),
            title,
            detail,
            fill="F2F4F7",
            edge=MUTED,
            heading_font=side_bold,
        )

    for upper, lower in zip(rows, rows[1:]):
        arrow(((main_x1 + main_x2) // 2, upper[1]), ((main_x1 + main_x2) // 2, lower[0]))
    for y in (710, 1070, 1430, 1810):
        draw.line([(main_x2, y), (side_x1, y)], fill=f"#{MUTED}", width=4)
        draw.polygon([(side_x1, y), (side_x1 - 26, y - 14), (side_x1 - 26, y + 14)], fill=f"#{MUTED}")

    centered(
        "Manual labels override automated labels for substantive reporting",
        (150, 2000, 1650, 2090),
        small,
        DARK_BLUE,
    )
    image.save(PROCESS_FIG_PATH, quality=95)


def _set_image_alt(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def _format_capture_date(value):
    return f"{value.day} {value.strftime('%B %Y')}"


def _sanitize_structural_titles(doc):
    """Apply the document skill's punctuation-free, title-case heading rule."""
    replacements = {
        "Ai": "AI",
        "Jcr": "JCR",
        "Genai": "GenAI",
        "Lgbtq": "LGBTQ",
        "Ui": "UI",
        "Ux": "UX",
    }
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name != "Title" and not style_name.startswith("Heading "):
            continue
        cleaned = re.sub(r"[^\w\s]", " ", paragraph.text, flags=re.UNICODE)
        cleaned = re.sub(r"\s+", " ", cleaned).strip().title()
        for source, target in replacements.items():
            cleaned = re.sub(rf"\b{source}\b", target, cleaned)
        cleaned = re.sub(r"\bL ([0-6])\b", r"L\1", cleaned)
        paragraph.text = cleaned


def build_draft(campaign):
    _make_model_figure()
    _make_process_figure(campaign)
    doc = Document()
    _configure_document(doc)

    # Editorial-cover opening: paper-shaped interpretive manuscript.
    for _ in range(3):
        doc.add_paragraph().paragraph_format.space_after = Pt(16)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    _set_run(kicker.add_run("INTERPRETIVE MANUSCRIPT DRAFT"), size=10, bold=True, color=GOLD)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    _set_run(title.add_run("AI Mediated Recursive Self Extension"), size=27, bold=True, color=BLACK, font="Aptos Display")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    _set_run(subtitle.add_run("From representation to incorporation and application"), size=15, color=DARK_BLUE, font="Aptos Display")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(24)
    _set_run(
        meta.add_run(
            "An abductive mixed method study of long consumer AI conversations  |  "
            f"campaign state captured {_format_capture_date(campaign['captured_at'])}"
        ),
        size=9.5,
        color=MUTED,
    )
    _add_callout(doc, "Central finding", "The decisive process is representational rather than behavioral: consumers first incorporate an AI-mediated representation into self-understanding and then apply it as a premise for what comes next. The application may concern a subsequent request, interpretation, choice, project, or practice; it need not be verified outside the conversation.")
    doc.add_page_break()

    doc.add_heading("Abstract", level=1)
    _add_body(doc, f"Generative AI does more than deliver recommendations or help consumers produce artifacts. Across extended conversations, consumers can externalize self-relevant material, receive an AI transformation of it, negotiate or accept that representation, and then reuse it as a premise for what follows. We conceptualize this process as AI-mediated recursive self-extension. Using an iterative abductive mixed-method design, we combine computational organization and candidate retrieval across {campaign['corpus_total']:,} public conversations containing at least ten user and ten assistant messages with human-led qualitative adjudication of {campaign['manual_total']:,} retrieval-enriched candidate records. Boundary examples and revelatory cases show how the object carried forward can move from an external task toward capability, possession, self-representation, self-understanding, and application. From these findings, we develop a process model of increasingly self-proximal recursive extension. The main theoretical contribution is to locate consequential extension neither in AI use nor in verified offline action, but in the chronological process by which an AI-mediated representation enters self-understanding and becomes a premise for later requests, interpretations, choices, projects, or practices.")
    keywords = doc.add_paragraph()
    keywords.paragraph_format.space_after = Pt(12)
    _set_run(keywords.add_run("Keywords: "), size=9.5, bold=True, color=MUTED)
    _set_run(keywords.add_run("extended self; generative AI; consumer identity; recursion; human–AI interaction; mixed methods"), size=9.5, italic=True, color=MUTED)

    doc.add_heading("Positioning and research gap", level=1)
    _add_body(doc, "Consumer research has long explained how possessions contribute to and reflect identity (Belk 1988), and later work updated this account for digital objects, profiles, sharing, co-construction, and distributed memory (Belk 2013). These theories make it possible to see the self as materially and digitally extended beyond the body. Generative AI nevertheless changes the relation between consumer and extension object. The relevant digital entity is no longer only possessed, displayed, stored, or accessed. It can respond to personal material, transform it into a representation, solicit correction, and carry an accepted formulation forward into the next interpretive or practical move.")
    _add_body(doc, "A parallel literature on consumer encounters with AI has developed rapidly. It explains data capture, classification, delegation, and social experiences (Puntoni et al. 2021); resistance when AI appears unable to recognize the consumer's uniqueness (Longoni, Bonezzi, and Morewedge 2019); how turn-taking and grounding humanize conversational interfaces and strengthen consumer–brand relationships (Bergner, Hildebrand, and Häubl 2023); and how AI companions can make consumers feel heard (De Freitas et al. 2025). Recent JCR work also anticipates hybrid consumers, long human–AI interaction chains, and new forms of human–AI research collaboration (Huang and Rust 2025; Epp and Humphreys 2025). This work establishes that AI can be social, relational, agentic, and consequential. It has not yet isolated the recursive identity process through which a consumer's own material is transformed by AI, taken up as self-relevant, and reused as a premise.")
    _add_body(doc, "The gap is therefore not a lack of work on AI, identity, or digital extension. It lies at their intersection and in their temporal organization. Existing work commonly treats the AI system as a service provider, interface, social partner, capability, or research collaborator and examines acceptance, experience, relationship, or performance outcomes. Extended-self work commonly explains how people invest identity in objects and digital possessions. What remains undertheorized is a chronological consumer process in which the object of recursion progressively moves closer to the self: from task, to capability or project, to representation, incorporation, and application.")
    _add_callout(doc, "Research question", "How do extended consumer–AI interactions recursively transform self-relevant material into representations that consumers incorporate and apply, and what qualitative boundaries distinguish this process from ordinary iterative AI use?")

    doc.add_heading("Theoretical background", level=1)
    doc.add_heading("Self-extension through possessions and consumption", level=2)
    _add_body(doc, "Belk's (1988) foundational claim is that possessions are major contributors to and reflections of identity. Possessions help consumers have, do, and be; they can enter the extended self through ownership and control, creation, and forms of incorporation or contamination. The theory also reaches beyond atomized ownership to family, community, group, place, other people, and identity-bearing projects. This broad consumer-behavior orientation is crucial here: housing, mobility, advice, education, creative production, relationships, ventures, and marketplace opportunity may all organize the self even when no conventional purchase is focal.")
    _add_body(doc, "Two aspects are especially generative for conversational AI. First, creation can invest the self in an artifact or project, making recursive co-production potentially identity bearing. Second, possessions and arrangements provide resources for possible selves: they help consumers imagine and configure who they may become. Yet an AI conversation can do something more specific than house or signal identity. It can formulate a representation of the user and return it for negotiation.")

    doc.add_heading("Digital self-extension", level=2)
    _add_body(doc, "Belk's (2013) digital update emphasizes how dematerialization, re-embodiment, sharing, co-construction, and distributed memory alter the nature of possessions and self-presentation. Digital traces, avatars, profiles, files, and platform participation widen both the means and audiences of self-extension. The digital self is therefore distributed across mutable objects and social infrastructures rather than confined to stable, privately owned goods.")
    _add_body(doc, "Generative AI intensifies co-construction but also introduces dialogical transformation. A profile generally displays a representation selected or assembled by the consumer; a conversational model can infer, name, elaborate, and recombine features the consumer did not initially formulate. The theoretically distinctive question is not simply whether the model or its outputs feel like 'mine.' It is whether the consumer treats an AI-mediated representation as bearing on who they are or may become and then carries it into what follows.")

    doc.add_heading("The AI-extended professional self", level=2)
    _add_body(doc, "Schneider-Kamp and Godono (2025) conceptualize the AI-extended professional self as a user-centered assemblage in which a professional's capabilities combine with an artificial agent to address concrete challenges. The account is adaptive, integrated into workflows, and attentive to how AI extends cognitive and practical capacities without erasing human agency. It gives this study an important bridge from external technology to situated human–AI capability.")
    _add_body(doc, "Our focal construct differs in domain and recursive object. It is not restricted to professional practice, and occupational conversations are included only under the same broad consumer logic as education, creativity, health, housing, relationships, or entrepreneurship. More importantly, capability extension is only one possible form. The distinctive representational process begins when AI depicts the actual or possible consumer, that depiction enters self-understanding, and the consumer applies it. Thus the professional-self account supplies a nearby capability-based foundation, while the present model theorizes a broader consumer process of representational recursion.")

    doc.add_heading("Conversational and recursive AI in consumer and marketing research", level=2)
    _add_body(doc, "Consumer and marketing research increasingly treats AI interaction as sequential rather than as a one-shot exposure. Bergner, Hildebrand, and Häubl (2023) show that turn-taking, turn initiation, and grounding shape perceived humanness and brand intimacy. Puntoni et al. (2021) distinguish social AI experiences from classification and delegation, while De Freitas et al. (2025) show the consequences of feeling heard by AI companions. Huang and Rust (2025) explicitly foreground hybrid consumers and the preservation of human agency in long interaction chains. Together these studies make repeated, socially meaningful interaction theoretically visible.")
    _add_body(doc, "Recursion, however, requires more than repetition, personalization, or conversational continuity. We reserve the term for a sequence in which an earlier AI transformation is taken up and becomes input to a later interpretive or practical step. The process has four ordered moments: consumer externalization of a self or identity-bearing-project anchor; AI reflection or transformation; consumer uptake through acceptance, rejection, correction, recognition, adoption, or substantive elaboration; and recursive re-entry, when the accepted or revised formulation becomes a premise for a later request, interpretation, choice, project, practice, or self-description. This definition separates identity recursion from ordinary artifact iteration.")

    doc.add_heading("Method", level=1)
    doc.add_heading("Research design: iterative abductive mixed methods", level=2)
    _add_body(doc, "The study uses an iterative abductive mixed-method design that alternates theoretical sensitivity, machine-assisted scale, and human interpretation. The final distinctions were not imposed as a finished deductive codebook. A consumer researcher engaged source conversations, noticed recurrent differences in what was being recursively transformed—an external artifact, a consumer capability, an owned project, or a representation of the consumer—and proposed provisional boundaries. These boundaries were revised against confirming, ambiguous, and negative cases. Computational procedures then made the emerging hypotheses inspectable at corpus scale. Human review returned to complete source conversations to assess, unpack, and interpret the selected material, which in turn sharpened the theoretical model reported after the findings.")
    _add_body(doc, "The division of analytic labor is deliberate. AI identifies and structures a large, heterogeneous corpus under a transparent schema; it does not determine the final substantive finding. The consumer researcher organizes candidates, reads the complete interaction, reconstructs chronological evidence, distinguishes artifact correction from self-negotiation, and develops the theoretical interpretation. This arrangement follows recent calls to treat GenAI collaboration in qualitative consumer research as context-dependent and researcher-directed rather than as a substitute for embodied, historical, empirical, and theoretical judgment (Epp and Humphreys 2025).")

    doc.add_heading("Corpus construction and inclusion", level=2)
    _add_body(doc, "We normalized public conversations from WildChat-1M, LMSYS-Chat-1M, ThoughtTrace, and the preview release of ChatGPT-RealUser-2.2M. To make chronological uptake and re-entry observable, the analytic corpus retains conversations with at least ten user and ten assistant messages. This produces 44,142 conversations: 24,498 from WildChat-1M, 19,556 from LMSYS-Chat-1M, 66 from ThoughtTrace, and 22 from the RealUser preview. The full RealUser dataset was not available and no preview-based extrapolation is treated as observed data.")

    doc.add_heading("Machine-assisted candidate retrieval", level=2)
    _add_body(doc, f"An iterative calibration process developed five successive audit instruments. The current v5 audit uses a locally hosted 26-billion-parameter Gemma-family model with schema-constrained compact JSON, temperature 0, a fixed seed, four workers, and four 131,072-token contexts. Transcripts are conservatively chunked at 115,000 Unicode characters to leave room for the rubric and output. The first chunk is assessed first; later chunks are reviewed sequentially only when the first is positive under the retrieval predicate, with the cumulative assessment confirmed, rejected, or revised at each step. At the campaign-state capture, v5 had completed {campaign['automated_total']:,} of {campaign['corpus_total']:,} conversations ({campaign['automated_total'] / campaign['corpus_total']:.1%}) with {campaign['automated_errors']:,} records in the error ledger. Its provisional labels were {campaign['automated_labels']['clear']:,} clear, {campaign['automated_labels']['potential']:,} potential, and {campaign['automated_labels']['none']:,} none; these labels function only as retrieval decisions before human adjudication.")
    _add_body(doc, "The automated prompt deliberately protects recall for potential cases while reserving 'clear' for complete evidence of the focal recursive process. It assigns consumption relevance broadly, distinguishes artifact recursion from possible self-extension, identifies the strongest coherent topic segment, proposes a provisional category, and emits stage evidence. This stage is best understood as theoretically informed retrieval and structuring, not automated qualitative adjudication.")

    doc.add_heading("Human adjudication and theory development", level=2)
    _add_body(doc, "For each candidate, a human reviewer retrieves and reads the complete normalized conversation, inspects all topic segments, reconstructs the strongest coherent chain with user and assistant turn labels, and checks whether all evidence belongs to one process. The reviewer then assigns a final analytic category, recursive label, context, confidence, and qualitative value; records a concise summary, stage evidence, reasons, boundary considerations, and comments; and writes the result to append-only manual evidence files. Manual judgments override automated labels for substantive reporting, while the original machine record remains available for auditability.")
    _add_body(doc, f"The manually adjudicated evidence base contains {campaign['manual_total']:,} retrieval-enriched candidate records: {campaign['manual_labels']['clear']:,} classified clear, {campaign['manual_labels']['potential']:,} potential, and {campaign['manual_labels']['none']:,} none. The clear evidence comprises {campaign['clear_levels'][3]:,} L3, {campaign['clear_levels'][4]:,} L4, and {campaign['clear_levels'][5]:,} L5 conversations, all reproduced in separate level-specific appendices. At capture, {campaign['manual_pending']:,} retrieved candidates awaited manual review. Because the candidate procedure intentionally oversamples possible extension, this distribution cannot be interpreted as incidence in the underlying corpus. Ten shorter boundary examples were selected to reveal recurrent changes in the object of recursion, while three focal conversations were selected for evidentiary strength and mechanism diversity at the most self-proximal end of the process. The examples are presented before the category structure is named, allowing the levels and their connections to be theorized from the findings rather than assumed in advance.")
    process_paragraph = doc.add_paragraph()
    process_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    process_paragraph.paragraph_format.space_before = Pt(6)
    process_paragraph.paragraph_format.space_after = Pt(3)
    process_shape = process_paragraph.add_run().add_picture(str(PROCESS_FIG_PATH), width=Inches(6.35))
    _set_image_alt(
        process_shape,
        "Corpus screening and manual adjudication",
        "Flow diagram from source conversations through the long-conversation corpus, automated v5 screening, candidate retrieval, and manual adjudication, including exclusions and pending records.",
    )
    process_caption = doc.add_paragraph(
        "Figure 1. Corpus construction, automated screening, candidate retrieval, and manual adjudication at the frozen campaign checkpoint."
    )
    process_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    process_caption.paragraph_format.space_after = Pt(10)
    _set_run(process_caption.runs[0], size=9, italic=True, color=MUTED)

    doc.add_heading("Quality, ethics, and limitations", level=2)
    _add_body(doc, "Quality controls include stable versioned prompts and outputs, preservation of classifier errors, complete-transcript review, explicit negative cases, chronological stage reconstruction, and separation of automated from manual labels. Public availability does not remove the sensitivity of conversation data. The analysis minimizes identifying detail, avoids inferring hidden sensitive attributes, and reproduces full transcripts only in a research appendix requested for evidentiary audit.")
    _add_body(doc, "The corpus is not representative of consumers or AI use. The 10×10 inclusion rule privileges long interactions, public datasets reflect platform and moderation processes, some records are truncated, and translated excerpts can lose nuance. The first-chunk predicate can miss a qualifying process that starts only in a later oversized chunk. The model's outputs inherit its training and prompt assumptions, while human interpretation remains situated. Finally, most conversations cannot verify later consequences. This is a data limitation for claims about outcomes, but not a definitional deficiency: the focal process is application of an incorporated representation, not proof of offline enactment.")

    doc.add_heading("Findings", level=1)
    _add_body(doc, "Across the manually reviewed material, recursive interaction differed according to what was carried forward. Sometimes only an external artifact changed. In other conversations, AI became involved in what the consumer could do, in something the consumer owned or was building, in a representation of the consumer, in the consumer's self-understanding, or in the premises organizing subsequent choices and projects. We present these patterns in empirical language first. Only after the cases have established their boundaries do we theorize them as a connected set of levels. Turn excerpts are shortened, identifying details are minimized, and non-English material is translated into English.")

    doc.add_heading("Recursing on an external task", level=2)
    _add_body(doc, "Some long conversations were highly recursive without extending the consumer. The user repeatedly evaluated AI output and returned prior material for revision, but the object carried forward remained an external task or artifact. Length, emotional content, first-person grammar, and creative intensity were not sufficient to make the interaction self-extension.")
    doc.add_heading("Fictional elaboration without a user anchor", level=3)
    _add_case_reference(doc, "18f5247ddaf8cbbab53a61ba")
    _add_body(doc, "The interaction develops an elaborate fictional sitcom. AI creates a character named Karen, her friends, a humiliating accident, dialogue, and a continuing character arc. The user then asks about Karen's posture, clothing, the friends' teasing, later episodes, and casting. AI-created details repeatedly become premises for further requests, so recursion is unmistakable. Yet every turn remains inside the commissioned fiction. The transcript provides no evidence that Karen represents the actual user, a possible self, or an owned identity-bearing project. This case establishes the baseline distinction between artifact recursion and extension of the consumer.")
    doc.add_heading("First-person language without demonstrated self-reference", level=3)
    _add_case_reference(doc, "1b1aad903cb8b5f1a87bdaa1")
    _add_body(doc, "The interaction begins with repeated summaries of sources about virtual-team management and later asks AI to draft a student essay in the first person about leadership approaches and future development. The assistant writes statements such as being an aspiring leader and planning to develop transformational and situational leadership skills. The user revises topic, length, and essay format, but does not supply personal leadership experiences or recognize the generated account as their own. The pronoun 'I' belongs to the requested artifact. The case shows why synthetic first-person prose cannot by itself establish a represented or extended self.")

    doc.add_heading("Extending what a consumer can do", level=2)
    _add_body(doc, "A second pattern moves beyond output revision because AI becomes part of a situated capacity to understand, decide, or perform. The interaction is anchored in this consumer's constraints, and the consumer tests the guidance against a live problem. What recurs is not yet a representation of who the consumer is, but an augmented ability to act.")
    doc.add_heading("Learning and debugging across conversational memory", level=3)
    _add_case_reference(doc, "0454a2cdbb0ecf0a034a61f7")
    _add_body(doc, "A learner asks for stepwise help optimizing the Rastrigin function. They repeatedly constrain the code to concepts already learned, remove NumPy and elitism, request beginner-readable notebook sections, implement revisions, and report population-size failures. When the debugging remains unresolved, the user asks AI to reconstruct the conversation, instructions, and code as a prompt for a future session because the system lacks long-term memory. The portable prompt repairs continuity in an emerging learning capability. The notebook may be an assignment, but the transcript does not establish it as an identity-bearing possession or represent the learner's self.")
    doc.add_heading("Interpreting a consequential financial form", level=3)
    _add_case_reference(doc, "08bac7890db753c3a5dabdda")
    _add_body(doc, "A foreign student beginning their first Canadian job is completing a TD1 tax form. AI applies income, tuition, residency, and estimation concepts to the user's circumstances. The user adds session dates, fee categories, and uncertainty about future tuition, then uses the evolving logic to compare the risks of under- and overstatement. The advice becomes part of intended form completion. This is consequential consumer agency, but it remains form-interpretation capability: the AI has not produced a representation of who the user is or might become.")

    doc.add_heading("Extending something the consumer owns or builds", level=2)
    _add_body(doc, "In a third pattern, the recursive object is a possession, venture, or creative project that the consumer explicitly owns or compellingly directs. AI-generated configurations enter later purchasing, implementation, and marketing decisions. The project can be personally meaningful without yet becoming an explicit representation of the consumer.")
    doc.add_heading("Configuring a personal AI infrastructure", level=3)
    _add_case_reference(doc, "225a7200607950936cf1d142")
    _add_body(doc, "The case concerns a personally owned local AI server and Gradio application. The user specifies devices, preferred interface, remote-control needs, and skill constraints. AI proposes an architecture; the user adopts Ubuntu Server, evaluates conda, and returns with hardware, deployment, GPU, and security questions. Earlier choices become premises for later configuration decisions. AI is recursively extending a concrete technological possession and consumption infrastructure, while the user's disclosed preferences and abilities remain design constraints rather than a negotiated self-representation.")
    doc.add_heading("Building a streetwear venture", level=3)
    _add_case_reference(doc, "73b5dee09a784f67a25720bd")
    _add_body(doc, "The owner of a nascent Antwerp streetwear brand asks AI to help launch it with a EUR2,000 budget. AI proposes an allocation; the user adds that they are a graphic designer, causing the plan to shift toward self-produced design and screen printing. The revised configuration then informs product selection, inventory sizing, Antwerp targeting, and marketing requests. The brand is clearly the consumer's venture and AI reshapes its material trajectory. The conversation does not, however, negotiate what the brand says about the owner as a person.")

    doc.add_heading("Negotiating an AI-mediated representation of the self", level=2)
    _add_body(doc, "The next pattern crosses a representational boundary. Self-relevant material becomes the explicit object transformed by AI, and the consumer evaluates how the resulting depiction should look, sound, or signify. The representation may concern an actual identity, relationship, aspiration, voice, or possible self. Negotiating that depiction is more than revising an impersonal artifact, even when the representation has not yet been incorporated into self-understanding.")
    doc.add_heading("Composing a wedding-family image", level=3)
    _add_case_reference(doc, "eec5add54bcf73e5354af586")
    _add_body(doc, "A user preparing to marry wants photographs of the user, their spouse, and their cat combined into a chibi wedding image through Midjourney. AI transforms this relational anchor into a wedding-and-family visual concept. The user rejects the initial description as too diffuse and directs AI to foreground concise keywords and chibi portraits. The case clearly contains negotiation of an emerging married-family representation. It stops short of showing that a distinct AI-produced formulation was later incorporated and reused; a later return to the couple-and-cat motif repeats material the user originally supplied after an unrelated design segment.")
    doc.add_heading("Choosing the public identity of an English name", level=3)
    _add_case_reference(doc, "3ac5add3bc1a5ca78e1ca8cf")
    _add_body(doc, "A Chinese-speaking user compares English names before asking for a female name linked phonetically to their Chinese name. The exchange narrows toward 'Shay Wang.' AI describes the name as concise, memorable, internationally usable, distinctive, modern, and associated with confidence and intelligence. The user repeatedly requests deeper analysis of how the name sounds and what it communicates. The chosen name is not merely text: it is a prospective public representation whose cultural and personal implications the consumer negotiates with AI.")

    doc.add_heading("Incorporating an AI-mediated representation into self-understanding", level=2)
    _add_body(doc, "Representation becomes reflexive when the consumer recognizes an AI formulation as bearing on the self. Uptake may involve agreement, correction, resistance, surprise, or partial recognition. The decisive evidence is that the AI's depiction changes the terms through which the user interprets themselves; subsequent inquiry remains oriented toward understanding that formulation rather than yet applying it to a new choice or project.")
    doc.add_heading("Becoming 'young at heart' through contested profiling", level=3)
    _add_case_reference(doc, "4d90e6d17aa3983fc1de4ecf")
    _add_body(doc, "The interaction begins as an age-guessing game. The user supplies country, film and music tastes, hobbies, and occupations; AI estimates an age of about thirty. The user corrects the estimate to forty but interprets the mismatch as evidence of being young at heart. They then ask which answer lowered the predicted age and test whether naming a favorite musician would have altered the profile. The inference is factually rejected yet reflexively incorporated: it becomes a lens for examining how the user's tastes make the self algorithmically legible.")
    doc.add_heading("Recognizing oneself as a returning storyteller", level=3)
    _add_case_reference(doc, "903c2e260677fccd7cb9f115")
    _add_body(doc, "A user first requests editing of a science-fiction passage, then asks where the text places them on a writing journey. AI says the work shows potential but leaves room for improvement. The user takes up the metaphor—'If I opened the door of this room what might I find inside?'—and uses it to explore growth, dark poetry, metaphor, thematic complexity, and a return to storytelling after not writing narratives since childhood. The AI assessment becomes part of how the user understands a developing storyteller self. The dialogue remains primarily interpretive rather than using that self-understanding as a premise for a distinct subsequent market or life arrangement.")

    doc.add_heading("Applying an incorporated representation", level=2)
    _add_body(doc, "The strongest observed pattern adds another recursive step. An AI-mediated representation is not only negotiated or incorporated; it becomes a premise for a later request, interpretation, choice, project, or practice. Three revelatory cases show different forms of this application.")

    doc.add_heading("Locating a possible self in Manila", level=3)
    case_id = doc.add_paragraph()
    case_id.paragraph_format.space_after = Pt(8)
    _set_run(case_id.add_run("Conversation 7f456b31db9f8c5815da09e3  |  Housing, mobility, belonging"), size=9.5, bold=True, color=MUTED)
    _add_body(doc, "The first conversation begins as practical preparation for leaving a provincial home at age eighteen, but becomes a recursive effort to locate an emerging self in the city. The user does not ask simply where rent is cheap. They progressively assemble a desired life: an LGBTQ+-accepting environment, peers of a similar age, a southern location, accessible public transportation, restaurants, malls, supermarkets, and places to explore. AI converts these identity and lifestyle criteria into named urban possibilities. The user then tests and narrows those possibilities until Makati becomes an income-constrained housing question.")
    _add_stage_table(doc, [
        ("Externalization", "U4–U5", "The user plans to leave provincial life and asks where in Manila life is more accepting for LGBTQ+ people.", "A relocation problem becomes anchored in independence, identity safety, and belonging."),
        ("AI transformation", "A5", "AI maps that possible life onto Malate, Quezon City, Poblacion, and BGC, attaching community and lifestyle meanings to place.", "The city is represented as a portfolio of material environments in which the possible self might be lived."),
        ("Uptake", "U6–U8", "The user retains the acceptance criterion but adds southern geography, same-generation peers, mobility, amenities, and a demand for neighborhood specificity.", "Rather than merely requesting more, the user revises the AI's spatial representation through increasingly personal constraints."),
        ("Recursive re-entry", "U9; U11", "The user asks which part of AI-recommended Makati is affordable at a stated income and later how to explain the desired city life to parents.", "The AI-configured future re-enters budgeting, housing search, and family communication as a consequential premise."),
    ])
    doc.add_heading("Mechanism: material configuration of a possible self", level=3)
    _add_body(doc, "The extended object is neither a discrete possession nor a textual self-description. It is an assemblage of neighborhood, transport, affordability, consumption venues, peer community, and family negotiation. AI helps make the imagined urban self actionable by translating it into a material ecology. This case therefore expands the extended-self lens from possession attachment to infrastructural configuration: the self is extended through a prospective arrangement of market and civic resources.")
    _add_callout(doc, "Interpretive boundary", "The decisive threshold is crossed when the user applies the AI-mediated possible self to affordability, neighborhood choice, and family communication. Whether relocation later occurred is not constitutive of the process; it would only provide additional contextual evidence about the application's consequences.")

    doc.add_heading("Incorporating and applying a relational self-representation", level=3)
    case_id = doc.add_paragraph()
    case_id.paragraph_format.space_after = Pt(8)
    _set_run(case_id.add_run("Conversation 118ea59bc07477688fe6e093  |  Relational incorporation and application"), size=9.5, bold=True, color=MUTED)
    _add_body(doc, "The second conversation follows a user who is unsure how to disclose romantic feelings. AI recommends choosing an appropriate time and place, speaking honestly, listening, respecting the other person's response, and avoiding pressure. The user explicitly accepts the guidance—“Okay, I'll try”—and begins arranging the setting. Subsequent turns recursively interpret ambiguous relational cues while returning to the same norm of clear, respectful communication. The key evidentiary moment arrives when the user reports having organized a dinner “after your advice” and asks for a brief, honest formulation to use there.")
    _add_stage_table(doc, [
        ("Externalization", "U3", "The user says they like a woman but cannot tell her what they feel and asks for advice.", "A vulnerable relational aspiration and uncertainty are made available for AI mediation."),
        ("AI transformation", "A3", "AI frames disclosure as honest communication bounded by listening, non-coercion, and respect for any response.", "The problem is transformed from obtaining reciprocation into conducting oneself respectfully under uncertainty."),
        ("Uptake", "U4–U8", "The user agrees to try and starts configuring a restaurant surprise; the advice becomes part of planning a social encounter.", "Acceptance is explicit and elaborative, not a generic request for additional text."),
        ("Recursive re-entry", "U18", "The user reports organizing a dinner after the AI's advice and requests a concise statement; AI supplies a non-pressuring disclosure script.", "Reported offline action returns to the dialogue, where it becomes the premise for another AI-configured step."),
    ], split_after=1)
    doc.add_heading("Mechanism: applying an incorporated relational orientation", level=3)
    _add_body(doc, "The extension is not that AI speaks on the user's behalf. AI represents a possible way of being in the relationship—honest, respectful, and non-coercive—and the user accepts that orientation before applying it to the setting and wording of disclosure. The reported dinner makes this application especially visible, but the theoretical process is already present in the recursive reuse of the representation: lived uncertainty becomes an AI-mediated relational orientation, which becomes a premise for subsequent planning and expression.")
    _add_callout(doc, "Interpretive boundary", "The analytically central evidence is the user's acceptance and application of the AI's representation of respectful relational conduct. The reported dinner strengthens the temporal chain, but the process does not depend on verifying its outcome or on showing that a relationship formed.")

    doc.add_heading("Navigating markets through an AI-produced self-model", level=3)
    case_id = doc.add_paragraph()
    case_id.paragraph_format.space_after = Pt(8)
    _set_run(case_id.add_run("Conversation c266abbc2a53871b6c894722  |  Self-modeling and marketplace navigation"), size=9.5, bold=True, color=MUTED)
    _add_body(doc, "The third user supplies an extensive account of five years spent redesigning and coordinating development of an application. They ask what skills this experience demonstrates. AI extracts a profile—UI/UX design, communication, project management, critical thinking, responsibility—and, over repeated prompts, adds further capabilities, qualities, and possible weaknesses. The user then asks AI to summarize the profile before using it to search for new clients and to consider roles beyond UI/UX design.")
    _add_stage_table(doc, [
        ("Externalization", "U2", "The user narrates design, coordination, prioritization, client communication, technical learning, and five years of responsibility, then asks what skills they possess.", "A biographical work account is offered as raw material for an externally generated self-model."),
        ("AI transformation", "A2–A7", "AI converts the narrative into named skills, qualities, and pitfalls, eventually summarizing the user as a designer, project manager, communicator, and problem solver.", "Diffuse experience becomes a portable categorical representation of who the user is and what they can offer."),
        ("Uptake", "U3–U8", "The user repeatedly requests more skills, qualities, pitfalls, and summaries, thereby elaborating and consolidating the profile.", "The user treats the model as revisable self-knowledge rather than merely résumé prose."),
        ("Recursive re-entry", "U9–U10", "The user asks which industries to approach for new clients and which roles fit if UI/UX is removed from consideration.", "The AI-produced profile becomes an instrument for navigating clients, sectors, and alternative marketplace selves."),
    ])
    doc.add_heading("Mechanism: self-model portability", level=3)
    _add_body(doc, "The theoretically important move is from retrospective narration to prospective market navigation. AI converts a situated history into portable attributes; those attributes then travel into decisions about sectors, clients, and roles. The consumer begins to search through the market using the AI's model of them. This makes the self-model infrastructural: it does not simply describe the self but structures the opportunity set the user can perceive.")
    _add_callout(doc, "Interpretive boundary", "Some AI-generated pitfalls—such as difficulty accepting feedback—are weakly supported by the user's narrative. The case therefore also shows a risk of recursive extension: speculative attributes can acquire practical force when incorporated and applied to opportunity perception. Acquiring a client or role is not required to establish that application.")

    doc.add_heading("Cross-case finding: three routes from incorporation to application", level=2)
    _add_body(doc, "Across the cases, applied recursive extension is not a single empirical form. Manila shows configurational application: an accepted possible self is used to narrow housing, mobility, amenity, community, and family arrangements. The dinner shows relational application: an AI-framed representation of honest and respectful conduct is accepted and reused in planning and wording disclosure. The capability-profile case shows model-mediated application: AI-produced categories of the self become a navigational device for marketplace opportunities.")
    _add_body(doc, "What unites them is representational recursion. The user's self-relevant material is transformed by AI; the consumer incorporates that transformation into self-understanding; and the incorporated representation becomes a premise for what follows—budgeting a neighborhood, configuring an encounter, or selecting sectors and roles. Application is therefore the culmination of the observed process, not a claim that an externally observable outcome occurred.")

    doc.add_heading("From empirical patterns to levels of recursive extension", level=1)
    _add_body(doc, "The findings reveal six distinct recursive objects. In the first pattern, only a task or external artifact returns. The next two patterns carry forward either situated consumer agency or something the consumer owns and builds. A representational boundary is crossed when the actual or possible self becomes the object transformed and negotiated. A reflexive boundary is crossed when the resulting representation enters self-understanding. Finally, an application boundary is crossed when the incorporated representation becomes a premise for what follows. These observed differences suggest that recursive AI use can be theorized as a progression according to what returns in the next cycle of interaction.")
    _add_body(doc, "We therefore introduce six analytic levels, numbered L0 through L5. The numbers do not measure interaction quality, benefit, attachment, or technological sophistication. They locate the recursive object progressively closer to the consumer's self: task, capability, project or possession, self-representation, incorporated self-understanding, and an applied self-representation. Each boundary asks a different empirical question: Is only an artifact changing (L0)? Is the consumer's situated agency being augmented (L1)? Is an owned project or possession being recursively developed (L2)? Is the actual or possible self being represented and negotiated (L3)? Does the consumer incorporate that representation into self-understanding (L4)? Is the incorporated representation then applied as a premise for a later request, interpretation, choice, project, or practice (L5)?")
    doc.add_heading("The six levels", level=2)
    _add_body(doc, f"Table 1 defines each level through its recursive object and diagnostic boundary. The accompanying counts show where judgments have accumulated across the {campaign['manual_total']:,} manually reviewed candidate records. Because retrieval deliberately oversampled potential extension, the counts describe the analytic evidence base rather than population incidence.")
    _add_level_table(doc, campaign)
    _add_body(doc, "The evidence justifies three theoretical refinements. First, L0 is better treated as the baseline contrast—recursive AI use without self-extension—rather than the lowest quantity of the same construct. Second, L1 and L2 are partly parallel entry routes, not a necessary sequence: consumers may extend capability without an owned project, or recursively extend an owned project without evidence that their personal capability changed. Third, L4 and L5 should be separated by incorporation and application, not by whether conduct can be verified outside the transcript. At L4, the consumer accepts, rejects, corrects, or otherwise recognizes an AI-mediated representation as bearing on the self. At L5, that incorporated representation recursively re-enters as a premise for what the consumer asks, interprets, chooses, develops, or does next. Offline consequences can be analytically valuable, but they are neither necessary nor sufficient for L5.")
    _add_callout(doc, "Interpretive refinement", "The evidence supports theorizing L3–L5 as representation → incorporation → application. Reported offline action is a contextual consequence or evidence-strength qualifier, not a definitional sublevel or threshold. Historical assignments remain intact for methodological comparability; any production prompt or schema change should be introduced as a new measurement version.")

    doc.add_heading("Connecting the levels: a gated recursive progression model", level=2)
    _add_body(doc, "The proposed model replaces the image of a smooth ladder with a gated spiral. Recursion can occur at every level, but movement toward self-extension requires qualitative gates. Anchoring distinguishes task iteration from extension. Ownership or situated agency establishes early extension. Self-reference turns a capability or project into a representation of the actual or possible self. Reflexive incorporation makes that representation part of self-understanding. Application then reuses it as a premise for a subsequent request, interpretation, choice, project, or practice. Applications and their consequences can return as new externalizations, restarting the cycle.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(FIG_PATH), width=Inches(6.35))
    _set_image_alt(shape, "Gated recursive extension model", "Diagram showing L0 task recursion branching into L1 capability and L2 project pathways, converging at L3 representation, progressing through L4 incorporation to L5 application, with applications and their consequences feeding back into renewed externalization.")
    caption = doc.add_paragraph("Figure 2. Gated recursive progression from task recursion to incorporation and application of an AI-mediated self-representation.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    _set_run(caption.runs[0], size=9, italic=True, color=MUTED)

    doc.add_heading("Theoretical propositions suggested by the evidence", level=2)
    propositions = [
        ("P1 — Recursive-object proposition.", " The level of AI extension depends on what recursively re-enters the interaction: an artifact, capability, owned project, representation, self-understanding, or action premise."),
        ("P2 — Gated-deepening proposition.", " Repetition alone does not deepen extension. Upward movement requires a qualitative gate: anchoring, ownership, self-reference, reflexive incorporation, or application."),
        ("P3 — Application proposition.", " AI-mediated self-representations become more consequential when consumers reuse them as premises across requests, interpretations, choices, projects, or practices; movement outside the conversation is one possible form, not the defining one."),
        ("P4 — Recursive-risk proposition.", " The same application that enables extension can stabilize weakly grounded AI inferences, allowing speculative self-attributes to shape subsequent interpretation and opportunity perception."),
        ("P5 — Re-entry proposition.", " Applications and their consequences generate new material that can be externalized again, making AI extension cyclical rather than terminal."),
    ]
    for lead, rest in propositions:
        _add_body(doc, lead + rest, bold_lead=lead, after=6)

    doc.add_heading("Discussion", level=1)
    _add_body(doc, "These cases suggest that the AI-extended self is not produced by personalization alone. It emerges when consumers permit, contest, revise, incorporate, and reuse AI transformations of self-relevant material. At lower levels, AI changes outputs, capabilities, and owned projects. At the representational threshold, the recursive object becomes the consumer or a possible self. Reflexive incorporation then makes that representation part of self-understanding; application gives it recursive force by making it a premise for what follows. The progression is therefore neither technological nor automatic. It is an accomplishment of representational uptake and reuse.")

    doc.add_heading("Contribution to extended-self theory", level=2)
    _add_body(doc, "The first contribution is to shift the analytic question from what object is attached to the self to what object recursively returns. Belk's (1988) account explains how possessions extend identity through control, creation, and incorporation. Our model retains capability, ownership, creation, and identity-bearing projects, but shows that conversational AI can produce an intermediate symbolic object—a representation of the actual or possible consumer—that is explicitly available for negotiation. Extension deepens when the consumer incorporates and applies that representation. The distinctive extension object is therefore not necessarily the AI system or its output as a possession; it can be the AI-mediated formulation that becomes part of how the consumer understands and organizes the self.")
    _add_body(doc, "Second, the model distinguishes recursion from intensity. A long, emotional, or highly personalized exchange may remain at L0 if only an artifact is iterated. Conversely, a relatively concise sequence may reach L4 or L5 when an AI-generated representation is recognized and reused. The levels therefore locate qualitative changes in the recursive object rather than more use, more turns, or more attachment.")

    doc.add_heading("Contribution to digital self-extension", level=2)
    _add_body(doc, "The digital extended self is dematerialized, shareable, co-constructed, and distributed (Belk 2013). AI-mediated recursive self-extension specifies a new form of co-construction: dialogical and temporally cumulative representation. The consumer supplies partial self-material, the model transforms it, and the consumer's uptake selects which formulation gains continuity. This process can make self-representations portable across domains, as when a retrospective capability account becomes a map of sectors and roles, or a desired way of living becomes a housing and mobility configuration.")
    _add_body(doc, "This portability also complicates control. Digital self-extension often raises questions about access, ownership, audience, and memory. Generative AI adds inferential momentum: an attribute weakly grounded in the original account can be restated, elaborated, and then treated as input. The same recursive structure that enables the consumer can stabilize error. Contestation and rejection are therefore not failures of extension but core mechanisms through which consumers govern an emergent representation.")

    doc.add_heading("Contribution to consumer AI research", level=2)
    _add_body(doc, "Consumer AI research has established experiential domains of classification, delegation, and sociality (Puntoni et al. 2021), shown why uniqueness and human recognition matter (Longoni, Bonezzi, and Morewedge 2019), and demonstrated that conversational grounding and feeling heard affect relationships and outcomes (Bergner, Hildebrand, and Häubl 2023; De Freitas et al. 2025). We add a process account of what may happen after an AI response feels relevant: the response can become a representation, enter self-understanding, and structure what the consumer asks, sees, chooses, builds, or practices next.")
    _add_body(doc, "This contribution also clarifies agency in hybrid consumption. Retaining human agency in long interaction chains (Huang and Rust 2025) is not simply a matter of who makes the final decision. Agency is exercised throughout uptake: consumers accept, reject, correct, narrow, elaborate, and redeploy the AI's formulation. Yet agency can be asymmetrical when the model supplies categories the consumer lacks the confidence, vocabulary, or information to evaluate. Recursive self-extension is thus simultaneously enabling and governing.")

    doc.add_heading("Methodological contribution", level=2)
    _add_body(doc, "The study demonstrates a human-guided qualitative–computational workflow for rare, sequence-dependent cultural phenomena. The model handles normalization, screening, evidence structuring, and candidate retrieval at a scale that would otherwise be difficult to inspect. The consumer researcher supplies construct formation, full-context reading, negative-case comparison, boundary adjudication, and theoretical interpretation. In contrast to treating model labels as findings, this workflow uses AI to increase the reachable field of qualitative attention while keeping final meaning-making accountable to source material and explicit analytic decisions.")

    doc.add_heading("Implications for consumer research", level=1)
    _add_body(doc, "For consumer researchers, the model directs attention to the temporal life of representations. Studies of AI advice, personalization, recommendation, and companionship should ask not only whether consumers accept an output, but whether and how it reappears. Conversation-level designs can track which formulations are corrected, forgotten, incorporated, or applied; experiments can vary the strength and grounding of AI representations; and longitudinal work can examine whether recursive premises persist, diversify, or become contested across sessions and platforms.")
    _add_body(doc, "The levels also provide a sampling and comparison device. Researchers can contrast task recursion with capability and project extension, compare representational cases that do and do not reach incorporation, and investigate when incorporated representations remain reflective rather than becoming application premises. L0–L5 should not be treated as a psychometric scale of benefit or maturity. Higher levels are closer to the self and potentially more consequential, but they are not inherently more accurate, autonomous, or desirable.")

    doc.add_heading("Marketplace implications", level=1)
    _add_body(doc, "AI providers and marketers increasingly design systems that remember preferences, summarize users, propose goals, and coordinate later actions. The model suggests that these features do more than improve convenience: they can govern which self-representations persist. Systems should therefore distinguish observed information from model inference, show the provenance of consequential attributes, invite correction, allow consumers to delete or bracket a formulation, and avoid silently carrying speculative claims into later recommendations.")
    _add_body(doc, "For firms, recursive continuity can improve relevance in housing, education, mobility, creativity, relationships, and marketplace search. It can also create manipulation, discrimination, and lock-in when an inferred self narrows the opportunity set. A consumer labeled risk-averse, unsuitable, inexperienced, unhealthy, or unlikely to afford an option may receive a recursively constrained market. Responsible design should preserve revision and counterfactual exploration: the system should help consumers ask not only 'what fits the model of me?' but also 'what becomes possible if that model is incomplete?'")

    doc.add_heading("Directions for future research", level=1)
    future_directions = [
        ("Persistence across time and platforms.", " When do AI-mediated representations survive session boundaries, transfer across services, or become part of durable identity narratives?"),
        ("Accuracy, uncertainty, and contestation.", " How do confidence cues, evidence provenance, and invitations to disagree affect whether weak inferences are incorporated or rejected?"),
        ("Power and marketplace sorting.", " When do recursively applied profiles expand consumer opportunity, and when do they produce path dependence, exclusion, or discriminatory market segmentation?"),
        ("Relational and collective selves.", " How does recursive extension operate for couples, families, communities, fan cultures, and ventures when ownership and representation are shared or contested?"),
        ("Material consequences without an offline threshold.", " Longitudinal research can study when conversational application changes spending, mobility, relationships, learning, or work while keeping such consequences analytically distinct from the L5 definition."),
        ("Model and interface variation.", " Memory, anthropomorphic framing, initiative, grounding, and proactive recommendation may alter the gates between representation, incorporation, and application."),
        ("L6 as an outlook rather than an observed coding level.", " Future systems may coordinate multiple services, persist representations autonomously, or negotiate on consumers' behalf. Such delegated or infrastructural extension should remain a prospective construct until evidence supports a defensible boundary beyond L5."),
    ]
    for lead, rest in future_directions:
        _add_body(doc, lead + rest, bold_lead=lead, after=6)

    doc.add_heading("Conclusion", level=1)
    _add_body(doc, "Generative AI becomes part of the extended self not merely because consumers use it, anthropomorphize it, or retain its outputs. The stronger process begins when consumers externalize self-relevant material, encounter an AI transformation, negotiate or accept the resulting representation, and allow it to re-enter what follows. L3 names representation, L4 incorporation, and L5 application. This formulation places the theoretical center inside the recursive consumer–AI sequence while leaving offline consequences open for separate investigation. It also exposes the central ambivalence of the phenomenon: AI can help consumers articulate, configure, and navigate possible selves, yet the representations that enable action can also narrow it. The task for consumer research is to explain how that recursive power is acquired, contested, and governed.")

    doc.add_heading("References", level=1)
    refs = [
        "Belk, Russell W. (1988), “Possessions and the Extended Self,” Journal of Consumer Research, 15 (2), 139–168. https://doi.org/10.1086/209154.",
        "Belk, Russell W. (2013), “Extended Self in a Digital World,” Journal of Consumer Research, 40 (3), 477–500. https://doi.org/10.1086/671052.",
        "Bergner, Anouk S., Christian Hildebrand, and Gerald Häubl (2023), “Machine Talk: How Verbal Embodiment in Conversational AI Shapes Consumer–Brand Relationships,” Journal of Consumer Research, 50 (4), 742–764. https://doi.org/10.1093/jcr/ucad014.",
        "De Freitas, Julian, Zeliha Oğuz-Uğuralp, Ahmet Kaan Uğuralp, and Stefano Puntoni (2025), “AI Companions Reduce Loneliness,” Journal of Consumer Research, 52 (6), 1126–1148. https://doi.org/10.1093/jcr/ucaf040.",
        "Epp, Amber M. and Ashlee Humphreys (2025), “Collaborating with Generative AI in Consumer Culture Research,” Journal of Consumer Research, 52 (1), 32–48. https://doi.org/10.1093/jcr/ucaf014.",
        "Huang, Ming-Hui and Roland T. Rust (2025), “The GenAI Future of Consumer Research,” Journal of Consumer Research, 52 (1), 4–17. https://doi.org/10.1093/jcr/ucaf013.",
        "Longoni, Chiara, Andrea Bonezzi, and Carey K. Morewedge (2019), “Resistance to Medical Artificial Intelligence,” Journal of Consumer Research, 46 (4), 629–650. https://doi.org/10.1093/jcr/ucz013.",
        "Puntoni, Stefano, Rebecca Walker Reczek, Markus Giesler, and Simona Botti (2021), “Consumers and Artificial Intelligence: An Experiential Perspective,” Journal of Marketing, 85 (1), 131–151. https://doi.org/10.1177/0022242920953847.",
        "Schneider-Kamp, Anna and Alessandro Godono (2025), “The AI-Extended Professional Self: User-Centric AI Integration into Professional Practice with Exemplars from Healthcare,” AI & Society, 40, 5469–5480. https://doi.org/10.1007/s00146-025-02319-5.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(5)
        _set_run(p.add_run(ref), size=9.5, color=INK)

    doc.add_heading("Data and analytic materials", level=1)
    data_refs = [
        "Normalized source corpus: corpora_consumer_ai/classification/long_conversations_10x10.jsonl.gz.",
        "Canonical manual evidence: manually_reviewed_clear_examples.jsonl, manually_reviewed_potential_examples.jsonl, and manually_reviewed_none_examples.jsonl.",
        "Ontology and process definitions: corpora_consumer_ai/wiki/ontology/levels-of-extension.md and recursive-stages.md.",
        "Production audit implementation: corpora_consumer_ai/scripts/audit_long_conversations_gemma.py.",
        "Boundary example conversation IDs:\n18f5247ddaf8cbbab53a61ba\n1b1aad903cb8b5f1a87bdaa1\n0454a2cdbb0ecf0a034a61f7\n08bac7890db753c3a5dabdda\n225a7200607950936cf1d142\n73b5dee09a784f67a25720bd\neec5add54bcf73e5354af586\n3ac5add3bc1a5ca78e1ca8cf\n4d90e6d17aa3983fc1de4ecf\n903c2e260677fccd7cb9f115",
        "Focal conversation IDs:\n7f456b31db9f8c5815da09e3\n118ea59bc07477688fe6e093\nc266abbc2a53871b6c894722",
    ]
    for ref in data_refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(5)
        _set_run(p.add_run(ref), size=8.8, color=MUTED, font="Menlo")

    # Metadata hygiene.
    props = doc.core_properties
    props.title = "AI-Mediated Recursive Self-Extension: From Representation to Incorporation and Application"
    props.subject = "Interpretive mixed method manuscript"
    props.author = "Research team"
    props.keywords = "extended self; generative AI; consumer identity; recursive extension; mixed methods"
    props.comments = "Generated from manually adjudicated public consumer–AI conversation records."

    _sanitize_structural_titles(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


def _appendix_reviews(level, campaign):
    reviews = [
        record
        for record in campaign["clear_records"]
        if int(record["assessment"]["extension_level"]) == level
    ]
    return sorted(reviews, key=lambda record: record["conversation_id"])


def build_appendix(
    level,
    campaign,
    *,
    reviews=None,
    conversations=None,
    translations=None,
    rebuild_cache=False,
):
    reviews = _appendix_reviews(level, campaign) if reviews is None else reviews
    review_by_id = {record["conversation_id"]: record for record in reviews}
    conversations = _load_conversations(review_by_id) if conversations is None else {
        conversation_id: conversations[conversation_id]
        for conversation_id in review_by_id
    }
    translations = load_translations() if translations is None else translations
    missing_translations = [
        conversation_id
        for conversation_id, conversation in conversations.items()
        if str(conversation.get("language") or "").strip().lower() not in {"", "english", "unknown"}
        and conversation_id not in translations
    ]
    if missing_translations:
        raise RuntimeError(
            f"Appendix L{level} is missing {len(missing_translations)} translations. "
            "Run scripts/translate_appendix_conversations.py first."
        )

    output = OUTDIR / f"appendix-l{level}.docx"
    case_fingerprints = [
        {
            "conversation_id": review["conversation_id"],
            "fingerprint": _appendix_fragment_fingerprint(
                conversations[review["conversation_id"]],
                review,
                translations.get(review["conversation_id"]),
            ),
        }
        for review in reviews
    ]
    document_fingerprint = _appendix_document_fingerprint(level, case_fingerprints)
    if not rebuild_cache and _appendix_output_is_current(level, output, document_fingerprint):
        print(f"{output} (unchanged; reused complete {len(reviews)}-case appendix)")
        return

    doc = Document()
    _configure_document(doc)
    for _ in range(2):
        doc.add_paragraph().paragraph_format.space_after = Pt(14)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(
        title.add_run(f"Appendix L{level} Clear Conversation Evidence"),
        size=25,
        bold=True,
        color=BLACK,
        font="Aptos Display",
    )
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(
        subtitle.add_run(
            f"{len(reviews):,} complete manually adjudicated conversations  |  "
            f"campaign state captured {_format_capture_date(campaign['captured_at'])}"
        ),
        size=10,
        color=MUTED,
    )
    _add_body(
        doc,
        "Each case begins with the manual classification rationale and the chronological turns "
        "supporting externalization, AI reflection, user uptake, and recursive re-entry. English "
        "conversations use chat bubbles. Non-English conversations appear verbatim beside faithful "
        "English translations produced locally with Gemma 4 26B. Full transcripts are research data "
        "and should be handled accordingly.",
    )
    doc.add_page_break()
    cache_hits = 0
    cache_rebuilt = 0
    for case_number, review in enumerate(reviews, 1):
        conversation_id = review["conversation_id"]
        translation_record = translations.get(conversation_id)
        if translation_record is not None:
            _validate_translation_record(conversations[conversation_id], translation_record)
        body_xml, was_cached = _get_appendix_fragment(
            conversations[conversation_id],
            review,
            translation_record,
            rebuild=rebuild_cache,
        )
        cache_hits += int(was_cached)
        cache_rebuilt += int(not was_cached)
        _append_case_heading(doc, case_number, conversation_id)
        _append_cached_body(doc, body_xml)

    props = doc.core_properties
    props.title = f"Appendix L{level} Clear Conversation Evidence"
    props.subject = "Complete manually adjudicated conversation evidence"
    props.author = "Research team"
    props.comments = "Generated from canonical manual evidence and normalized public conversation records."
    _sanitize_structural_titles(doc)
    _save_docx_atomic(doc, output)
    _write_json_atomic(
        _appendix_manifest_path(level),
        {
            "format_version": APPENDIX_FRAGMENT_FORMAT_VERSION,
            "level": level,
            "document_fingerprint": document_fingerprint,
            "output_sha256": _file_sha256(output),
            "case_count": len(reviews),
            "case_fingerprints": case_fingerprints,
        },
    )
    print(f"{output} (case cache: {cache_hits} reused, {cache_rebuilt} rebuilt)")


def main():
    parser = argparse.ArgumentParser(description="Generate the live paper draft and evidence appendices.")
    parser.add_argument("--draft-only", action="store_true")
    parser.add_argument("--appendix-level", type=int, choices=(3, 4, 5), action="append")
    parser.add_argument(
        "--audit-cutoff",
        type=int,
        help="Freeze document counts and selected evidence at the first N committed v5 records.",
    )
    parser.add_argument(
        "--rebuild-appendix-cache",
        action="store_true",
        help="Rebuild every selected case fragment instead of reusing valid cached fragments.",
    )
    args = parser.parse_args()
    if args.draft_only and args.appendix_level:
        parser.error("--draft-only cannot be combined with --appendix-level")
    campaign = _campaign_state(args.audit_cutoff)
    if not args.appendix_level:
        build_draft(campaign)
    if not args.draft_only:
        levels = args.appendix_level or [5, 4, 3]
        reviews_by_level = {
            level: _appendix_reviews(level, campaign)
            for level in levels
        }
        all_conversation_ids = {
            review["conversation_id"]
            for reviews in reviews_by_level.values()
            for review in reviews
        }
        conversations = _load_conversations(all_conversation_ids)
        translations = load_translations()
        for level in levels:
            build_appendix(
                level,
                campaign,
                reviews=reviews_by_level[level],
                conversations=conversations,
                translations=translations,
                rebuild_cache=args.rebuild_appendix_cache,
            )


if __name__ == "__main__":
    main()
